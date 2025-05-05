

from googleapiclient.discovery import build
import openpyxl
from openpyxl.utils import column_index_from_string
import os

YOUR_API_KEY = "AIzaSyCCDcqtKPJE46_IiHC-UHz6HHvwQ5muoog"
YOUTUBE_VIDEO_ID = "gGSwlyCYe58"
SAVE_PATH = r"C:\Users\charlton\Desktop"
EXCEL_FILENAME = f"{YOUTUBE_VIDEO_ID}_all.xlsx"
FULL_PATH = os.path.join(SAVE_PATH, EXCEL_FILENAME)

try:
    youtube = build("youtube", "v3", developerKey=YOUR_API_KEY)

    comments = []
    next_page_token = None

    print("전체 댓글 데이터를 수집 중입니다...")

    while True:
        request = youtube.commentThreads().list(
            part="snippet",
            videoId=YOUTUBE_VIDEO_ID,
            maxResults=100,
            pageToken=next_page_token
        )
        response = request.execute()

        comments_data = response.get("items", [])
        if not comments_data:
            break

        comments.extend([{'text': item['snippet']['topLevelComment']['snippet']['textDisplay'],
                           'like_count': item['snippet']['topLevelComment']['snippet'].get('likeCount', 0)}
                         for item in comments_data])

        print(f"이번 페이지에서 {len(comments_data)}개의 댓글 수집됨. 총 {len(comments)}개.")

        next_page_token = response.get("nextPageToken")
        if next_page_token is None:
            print("더 이상 다음 페이지가 없습니다.")
            break
        # 최대 수집 개수 제한을 제거하거나 늘립니다.
        # if len(comments) >= 1000:
        #     print("최대 수집 개수에 도달하여 중단합니다.")
        #     break

    if not comments:
        print("댓글이 없습니다.")
    else:
        # 좋아요 수 기준으로 내림차순 정렬
        sorted_comments = sorted(comments, key=lambda x: x['like_count'], reverse=True)

        # 엑셀 파일 생성
        workbook = openpyxl.Workbook()
        sheet = workbook.active

        # 헤더 행 추가
        sheet['A1'] = "Index"
        sheet['B1'] = "내용"
        sheet['C1'] = "좋아요 수"

        print("엑셀 파일에 데이터를 쓰는 중입니다...")

        # 댓글 데이터 추가
        for i, comment in enumerate(sorted_comments):
            row_index = i + 2
            sheet.cell(row=row_index, column=column_index_from_string('A'), value=i + 1)
            sheet.cell(row=row_index, column=column_index_from_string('B'), value=comment['text'])
            sheet.cell(row=row_index, column=column_index_from_string('C'), value=comment['like_count'])

        # 엑셀 파일 저장
        workbook.save(FULL_PATH)
        print(f"전체 댓글 데이터를 '{FULL_PATH}'에 저장했습니다. (총 {len(sorted_comments)}개)")

except Exception as e:
    print(f"오류 발생: {e}")